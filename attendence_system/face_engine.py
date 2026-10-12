"""
face_engine.py - shared detection / alignment / recognition / quality utilities.

Models (InsightFace model packs, run with ONNX Runtime):
  * Detection  : SCRFD-10G  (det_10g.onnx)   - strong on small / occluded faces
  * Recognition: ArcFace R50 trained on WebFace600K (w600k_r50.onnx) -> 512-d embedding
  ("antelopev2" pack = SCRFD-10G + ArcFace R100 on Glint360K: more accurate, ~2x slower)
"""
from __future__ import annotations

import logging
import os
from dataclasses import dataclass, field

import cv2
import numpy as np
from PIL import Image, ImageOps

try:  # iPhone photos are often HEIC
    from pillow_heif import register_heif_opener

    register_heif_opener()
except ImportError:
    pass

from insightface.app import FaceAnalysis
from insightface.utils import face_align

log = logging.getLogger("face_engine")

VIEWS = ("front", "left45", "right45")


# ----------------------------------------------------------------------------- I/O
def load_image(path: str, max_side: int | None = None) -> np.ndarray:
    """Load jpg/png/webp/heic as BGR. Applies EXIF rotation - phone photos are often
    stored sideways with a rotation flag, and an un-rotated face is a common cause of
    'no face detected'."""
    with Image.open(path) as im:
        im = ImageOps.exif_transpose(im).convert("RGB")
        if max_side and max(im.size) > max_side:
            im.thumbnail((max_side, max_side), Image.LANCZOS)
        arr = np.asarray(im)
    return cv2.cvtColor(arr, cv2.COLOR_RGB2BGR)


def save_jpg(path: str, img: np.ndarray, quality: int = 92) -> str:
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    cv2.imwrite(path, img, [cv2.IMWRITE_JPEG_QUALITY, quality])
    return path


def crop_with_margin(img: np.ndarray, bbox, margin: float = 0.3) -> np.ndarray:
    h, w = img.shape[:2]
    x1, y1, x2, y2 = bbox[:4]
    bw, bh = x2 - x1, y2 - y1
    x1 = int(max(0, x1 - margin * bw)); y1 = int(max(0, y1 - margin * bh))
    x2 = int(min(w, x2 + margin * bw)); y2 = int(min(h, y2 + margin * bh))
    return img[y1:y2, x1:x2].copy()


# ------------------------------------------------------------------------- quality
def blur_score(aligned: np.ndarray) -> float:
    """Variance of Laplacian on the 112x112 aligned crop. Because every face is resampled
    to the same size, this captures both motion/focus blur AND low resolution (a 20px face
    upsampled to 112px is smooth -> low score)."""
    gray = cv2.cvtColor(aligned, cv2.COLOR_BGR2GRAY)
    return float(cv2.Laplacian(gray, cv2.CV_64F).var())


def yaw_ratio(kps: np.ndarray) -> float:
    """Heuristic yaw from 5 landmarks: horizontal nose offset from eye midpoint, divided by
    eye distance. ~0 frontal, ~0.3-0.5 at 45 deg, >0.8 near profile (rough, uncalibrated)."""
    le, re, nose = kps[0], kps[1], kps[2]
    eye_dist = max(np.linalg.norm(re - le), 1.0)
    return float(abs(nose[0] - (le[0] + re[0]) / 2.0) / eye_dist)


@dataclass
class QualityConfig:
    min_face_px: int = 24        # min(bbox w, h) in original pixels
    min_blur: float = 25.0       # Laplacian variance on aligned crop
    min_det_score: float = 0.5
    max_yaw: float = 0.9         # yaw_ratio above this ~ profile face
    min_feat_norm: float = 0.0   # raw ArcFace norm (low norm ~ low quality); 0 = disabled


@dataclass(eq=False)
class Face:
    bbox: np.ndarray
    score: float
    kps: np.ndarray
    aligned: np.ndarray | None = None
    embedding: np.ndarray | None = None
    feat_norm: float = 0.0
    face_px: int = 0
    blur: float = 0.0
    yaw: float = 0.0
    reject_reasons: list = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.reject_reasons


def assess(face: Face, q: QualityConfig) -> None:
    r = face.reject_reasons
    if face.face_px < q.min_face_px: r.append(f"small({face.face_px}px)")
    if face.blur < q.min_blur: r.append(f"blurry({face.blur:.0f})")
    if face.score < q.min_det_score: r.append(f"low_det({face.score:.2f})")
    if face.yaw > q.max_yaw: r.append(f"profile({face.yaw:.2f})")
    if q.min_feat_norm and face.feat_norm < q.min_feat_norm:
        r.append(f"low_norm({face.feat_norm:.1f})")


# ----------------------------------------------------------------------------- NMS
def nms(boxes: np.ndarray, scores: np.ndarray, iou_thr: float = 0.4) -> list[int]:
    x1, y1, x2, y2 = boxes.T
    areas = (x2 - x1) * (y2 - y1)
    order = scores.argsort()[::-1]
    keep = []
    while order.size:
        i = order[0]; keep.append(int(i))
        xx1 = np.maximum(x1[i], x1[order[1:]]); yy1 = np.maximum(y1[i], y1[order[1:]])
        xx2 = np.minimum(x2[i], x2[order[1:]]); yy2 = np.minimum(y2[i], y2[order[1:]])
        inter = np.clip(xx2 - xx1, 0, None) * np.clip(yy2 - yy1, 0, None)
        iou = inter / (areas[i] + areas[order[1:]] - inter + 1e-9)
        order = order[1:][iou <= iou_thr]
    return keep


# --------------------------------------------------------------------------- engine
class FaceEngine:
    def __init__(self, model_pack: str = "buffalo_l", det_size: int = 640,
                 det_thresh: float = 0.5, use_gpu: bool = False,
                 model_root: str = "~/.insightface"):
        providers = (["CUDAExecutionProvider", "CPUExecutionProvider"] if use_gpu
                     else ["CPUExecutionProvider"])
        self.app = FaceAnalysis(name=model_pack, root=os.path.expanduser(model_root),
                                allowed_modules=["detection", "recognition"],
                                providers=providers)
        self.app.prepare(ctx_id=0 if use_gpu else -1, det_thresh=det_thresh,
                         det_size=(det_size, det_size))
        self.det = self.app.det_model
        self.rec = self.app.models["recognition"]
        self.det_size = det_size
        self.model_pack = model_pack
        log.info("Loaded %s (det_size=%d, providers=%s)", model_pack, det_size, providers)

    # -- detection -----------------------------------------------------------------
    def _detect_once(self, img, size):
        size = int(round(size / 32) * 32)
        boxes, kps = self.det.detect(img, input_size=(size, size), max_num=0)
        if kps is None:
            kps = np.zeros((len(boxes), 5, 2), np.float32)
        return boxes, kps

    def detect(self, img: np.ndarray, det_size: int | None = None, tiles: int = 1,
               tile_overlap: float = 0.25) -> list[Face]:
        """Full-image pass + optional tiles x tiles overlapping grid pass, merged by NMS.
        Root cause it fixes: SCRFD resizes the whole photo to det_size, so a 40px face
        in a 4000px class photo shrinks to ~13px at 1280 and gets missed. Tiles keep
        each crop closer to native resolution."""
        size = det_size or self.det_size
        all_b, all_k = [], []
        b, k = self._detect_once(img, size)
        all_b.append(b); all_k.append(k)

        if tiles > 1:
            H, W = img.shape[:2]
            th = int(min(H, H / tiles * (1 + tile_overlap)))
            tw = int(min(W, W / tiles * (1 + tile_overlap)))
            ys = np.linspace(0, H - th, tiles).astype(int)
            xs = np.linspace(0, W - tw, tiles).astype(int)
            for y0 in ys:
                for x0 in xs:
                    tb, tk = self._detect_once(img[y0:y0 + th, x0:x0 + tw], size)
                    if not len(tb):
                        continue
                    # drop faces cut by an *inner* tile border (partial faces)
                    m = np.ones(len(tb), bool)
                    if x0 > 0: m &= tb[:, 0] > 2
                    if y0 > 0: m &= tb[:, 1] > 2
                    if x0 + tw < W: m &= tb[:, 2] < tw - 2
                    if y0 + th < H: m &= tb[:, 3] < th - 2
                    tb, tk = tb[m].copy(), tk[m].copy()
                    tb[:, [0, 2]] += x0; tb[:, [1, 3]] += y0
                    tk[:, :, 0] += x0; tk[:, :, 1] += y0
                    all_b.append(tb); all_k.append(tk)

        boxes = np.concatenate(all_b) if all_b else np.zeros((0, 5))
        kpss = np.concatenate(all_k) if all_k else np.zeros((0, 5, 2))
        if not len(boxes):
            return []
        keep = nms(boxes[:, :4], boxes[:, 4]) if tiles > 1 else range(len(boxes))
        faces = []
        for i in keep:
            bb = boxes[i]
            faces.append(Face(bbox=bb[:4].astype(np.float32), score=float(bb[4]),
                              kps=kpss[i].astype(np.float32),
                              face_px=int(min(bb[2] - bb[0], bb[3] - bb[1]))))
        return faces

    # -- alignment + quality ---------------------------------------------------------
    def align(self, img: np.ndarray, faces: list[Face]) -> None:
        for f in faces:
            f.aligned = face_align.norm_crop(img, landmark=f.kps, image_size=112)
            f.blur = blur_score(f.aligned)
            f.yaw = yaw_ratio(f.kps)

    # -- recognition -----------------------------------------------------------------
    def embed(self, faces: list[Face], flip_tta: bool = True, batch: int = 64) -> None:
        """Batch ArcFace. flip_tta averages the embedding of the crop and its mirror,
        which makes side faces more stable (a 45-deg-left face looks like 45-deg-right)."""
        todo = [f for f in faces if f.aligned is not None]
        for s in range(0, len(todo), batch):
            chunk = todo[s:s + batch]
            imgs = [f.aligned for f in chunk]
            if flip_tta:
                imgs += [cv2.flip(a, 1) for a in imgs]
            feats = self.rec.get_feat(imgs)
            n = len(chunk)
            raw = feats[:n] + feats[n:] if flip_tta else feats
            norms = np.linalg.norm(feats[:n], axis=1)
            emb = raw / np.linalg.norm(raw, axis=1, keepdims=True)
            for f, e, nm in zip(chunk, emb, norms):
                f.embedding = e.astype(np.float32)
                f.feat_norm = float(nm)


# ------------------------------------------------------------------------ excel util
def make_thumb(src: str | np.ndarray, dst: str, max_px: int = 110) -> str:
    img = load_image(src) if isinstance(src, str) else src
    h, w = img.shape[:2]
    s = max_px / max(h, w)
    if s < 1:
        img = cv2.resize(img, (int(w * s), int(h * s)), interpolation=cv2.INTER_AREA)
    return save_jpg(dst, img, 85)


def write_sheet(ws, headers, rows, image_cols=(), row_height=88, col_width=16):
    """Write a table; columns listed in image_cols hold a thumbnail path that is
    embedded as a picture instead of text."""
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    img_idx = {headers.index(h) for h in image_cols}
    for r, row in enumerate(rows, start=2):
        for ci, val in enumerate(row):
            letter = get_column_letter(ci + 1)
            if ci in img_idx:
                if val and os.path.exists(val):
                    ws.add_image(XLImage(val), f"{letter}{r}")
                continue
            ws[f"{letter}{r}"] = val
            ws[f"{letter}{r}"].alignment = Alignment(vertical="center", wrap_text=True)
        if img_idx:
            ws.row_dimensions[r].height = row_height
    for ci, h in enumerate(headers):
        ws.column_dimensions[get_column_letter(ci + 1)].width = 17 if ci in img_idx else col_width
    ws.freeze_panes = "A2"
