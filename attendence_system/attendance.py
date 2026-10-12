"""
attendance.py - classroom photos -> detect -> quality gate -> embed -> match -> attendance.

Usage:
    python attendance.py --images class_photos/ --db enroll_output/embeddings.pkl --out run_2026_10_01
    python attendance.py --images a.jpg b.jpg c.jpg --db enroll_output/embeddings.pkl --gpu

Outputs (in --out):
    attendance.xlsx          one row per enrolled student: Present / Review / Absent
    validation_report.xlsx   one row per detected face: crop | matched enroll crop | enroll Image1
                             + scores, plus a 'Rejected' sheet for faces dropped by quality checks
    annotated/               class photos with boxes and predicted roll numbers
    faces/                   every detected face crop
"""
from __future__ import annotations

import argparse
import glob
import logging
import os
import pickle
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from scipy.optimize import linear_sum_assignment
from tqdm import tqdm

from face_engine import (FaceEngine, QualityConfig, assess, crop_with_margin, load_image,
                         make_thumb, save_jpg, write_sheet)

log = logging.getLogger("attendance")
IMG_EXT = (".jpg", ".jpeg", ".png", ".webp", ".bmp", ".heic", ".heif")


# ------------------------------------------------------------------------- gallery
def build_gallery(db):
    """Templates = each enrolled view + the per-student mean. Score for a student is the
    max over its templates, so a side-face in class can match the 45-deg enroll image."""
    G, owner, view = [], [], []
    for si, s in enumerate(db["students"]):
        for vi, v in enumerate(s["views"]):
            G.append(s["embeddings"][vi]); owner.append(si); view.append(v)
        if s["mean_embedding"] is not None:
            G.append(s["mean_embedding"]); owner.append(si); view.append("mean")
    if not G:
        raise SystemExit("No enrolled embeddings in DB.")
    return np.stack(G).astype(np.float32), np.array(owner), np.array(view)


def student_scores(E, G, owner, n_students):
    """E (F,512) x G (T,512) -> per-student max similarity (F,S) and argmax template (F,S)."""
    sim = E @ G.T
    S = np.full((len(E), n_students), -1.0, np.float32)
    T = np.full((len(E), n_students), -1, np.int64)
    for si in range(n_students):
        idx = np.where(owner == si)[0]
        if len(idx):
            j = sim[:, idx].argmax(1)
            S[:, si] = sim[np.arange(len(E)), idx[j]]
            T[:, si] = idx[j]
    return S, T


def assign(S, min_score):
    """One-to-one face->student within ONE photo (a student can't appear twice in it).
    Hungarian maximises total similarity, which resolves look-alike conflicts."""
    rows, cols = linear_sum_assignment(-S)
    return {int(r): int(c) for r, c in zip(rows, cols) if S[r, c] >= min_score}


# ---------------------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description="Classroom attendance")
    ap.add_argument("--images", nargs="+", required=True, help="folder(s) and/or image files")
    ap.add_argument("--db", default="enroll_output/embeddings.pkl")
    ap.add_argument("--out", default=f"attendance_{datetime.now():%Y%m%d_%H%M}")
    ap.add_argument("--gpu", action="store_true")
    ap.add_argument("--det-size", type=int, default=1280)
    ap.add_argument("--tiles", type=int, default=2, help="NxN tiled detection (1 = off)")
    ap.add_argument("--det-thresh", type=float, default=0.45)
    ap.add_argument("--match-thresh", type=float, default=0.40, help="cosine sim => Present")
    ap.add_argument("--review-thresh", type=float, default=0.30, help="cosine sim => Review")
    ap.add_argument("--min-face", type=int, default=24)
    ap.add_argument("--min-blur", type=float, default=25.0)
    ap.add_argument("--max-yaw", type=float, default=0.9)
    ap.add_argument("--max-side", type=int, default=4000, help="downscale huge photos")
    args = ap.parse_args()

    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    out = args.out
    for d in ("annotated", "faces", "thumbs"):
        os.makedirs(os.path.join(out, d), exist_ok=True)

    with open(args.db, "rb") as fh:
        db = pickle.load(fh)
    students = db["students"]
    G, owner, tview = build_gallery(db)
    q = QualityConfig(min_face_px=args.min_face, min_blur=args.min_blur, max_yaw=args.max_yaw,
                      min_det_score=args.det_thresh)

    paths = []
    for p in args.images:
        if os.path.isdir(p):
            paths += sorted(f for f in glob.glob(os.path.join(p, "*")) if f.lower().endswith(IMG_EXT))
        else:
            paths.append(p)
    if not paths:
        raise SystemExit("No class images found.")

    engine = FaceEngine(db.get("model_pack", "buffalo_l"), det_size=args.det_size,
                        det_thresh=args.det_thresh, use_gpu=args.gpu)
    flip = db.get("flip_tta", True)

    face_rows, rej_rows = [], []
    best = {}          # student idx -> dict(best evidence)
    seen_in = {}       # student idx -> set(image names) where matched >= match_thresh

    def _load(p):
        try:
            return p, load_image(p, max_side=args.max_side), None
        except Exception as e:  # noqa: BLE001
            return p, None, e

    with ThreadPoolExecutor(max_workers=4) as ex:   # decode next photos while GPU/CPU infers
        for p, img, err in tqdm(ex.map(_load, paths), total=len(paths), desc="Class photos"):
            name = os.path.basename(p)
            if err is not None:
                log.warning("skip %s: %s", p, err); continue
            stem = os.path.splitext(name)[0]

            faces = engine.detect(img, tiles=args.tiles)
            engine.align(img, faces)
            engine.embed(faces, flip_tta=flip)       # embed all; norm is also a quality signal
            for f in faces:
                assess(f, q)
            good = [f for f in faces if f.ok]
            log.info("%s: %d faces detected, %d passed quality", name, len(faces), len(good))

            gidx = {id(f): i for i, f in enumerate(good)}
            pred = {}
            S = T = None
            if good:
                E = np.stack([f.embedding for f in good])
                S, T = student_scores(E, G, owner, len(students))
                pred = assign(S, args.review_thresh)

            vis = img.copy()
            for fi, f in enumerate(faces):
                crop_p = save_jpg(os.path.join(out, "faces", f"{stem}_f{fi:03d}.jpg"),
                                  crop_with_margin(img, f.bbox, 0.2))
                x1, y1, x2, y2 = f.bbox.astype(int)
                if not f.ok:
                    rej_rows.append([name, fi, make_thumb(crop_p, os.path.join(out, "thumbs", f"r_{stem}_{fi}.jpg")),
                                     ", ".join(f.reject_reasons), f.face_px, round(f.blur, 1),
                                     round(f.score, 3), round(f.yaw, 2), round(f.feat_norm, 1)])
                    cv2.rectangle(vis, (x1, y1), (x2, y2), (128, 128, 128), 2)
                    continue

                gi = gidx[id(f)]
                order = np.argsort(-S[gi])
                top1, top2 = float(S[gi, order[0]]), float(S[gi, order[1]]) if len(order) > 1 else -1.0
                si = pred.get(gi)
                if si is None:
                    decision, score, stu, matched_crop = "UNKNOWN", top1, None, None
                else:
                    score = float(S[gi, si]); stu = students[si]
                    decision = "MATCH" if score >= args.match_thresh else "REVIEW"
                    v = tview[T[gi, si]]
                    v = stu["views"][0] if v == "mean" else v
                    matched_crop = stu["crops"].get(v)
                    if decision == "MATCH":
                        seen_in.setdefault(si, set()).add(name)
                    if si not in best or score > best[si]["score"]:
                        best[si] = {"score": score, "image": name, "crop": crop_p}

                color = {"MATCH": (0, 200, 0), "REVIEW": (0, 165, 255), "UNKNOWN": (0, 0, 255)}[decision]
                cv2.rectangle(vis, (x1, y1), (x2, y2), color, 2)
                label = f"{stu['roll'] if stu else '?'} {score:.2f}"
                cv2.putText(vis, label, (x1, max(12, y1 - 5)), cv2.FONT_HERSHEY_SIMPLEX,
                            max(0.4, (x2 - x1) / 120), color, 2)

                th = lambda src, tag: make_thumb(src, os.path.join(out, "thumbs", f"{tag}_{stem}_{fi}.jpg")) if src else None
                face_rows.append([
                    name, fi, th(crop_p, "d"),
                    th(matched_crop, "m"),
                    th(stu["front_source"], "o") if stu and stu.get("front_source") else None,
                    stu["roll"] if stu else "", stu["name"] if stu else "",
                    round(score, 3), round(top1 - top2, 3), decision,
                    students[order[0]]["roll"], f.face_px, round(f.blur, 1), round(f.score, 3),
                    round(f.yaw, 2), round(f.feat_norm, 1)])
            save_jpg(os.path.join(out, "annotated", f"{stem}_annotated.jpg"), vis, 88)

    # ---------------------------------------------------------- attendance sheet
    wb = Workbook(); ws = wb.active; ws.title = "Attendance"
    date = datetime.now().strftime("%Y-%m-%d %H:%M")
    rows = []
    for si, s in enumerate(students):
        if not s["views"]:
            status, sc, imgs, ev = "NOT ENROLLED", None, 0, ""
        elif si in best:
            sc = round(best[si]["score"], 3)
            status = "Present" if sc >= args.match_thresh else "Review"
            imgs, ev = len(seen_in.get(si, ())), best[si]["image"]
        else:
            status, sc, imgs, ev = "Absent", None, 0, ""
        rows.append([s["roll"], s["name"], status, sc, imgs, ev, date])
    write_sheet(ws, ["Roll No", "Name", "Status", "Best similarity", "# photos matched",
                     "Best evidence photo", "Taken at"], rows, col_width=20)
    fills = {"Present": "C6EFCE", "Absent": "FFC7CE", "Review": "FFEB9C", "NOT ENROLLED": "D9D9D9"}
    for r in range(2, ws.max_row + 1):
        ws[f"C{r}"].fill = PatternFill("solid", fgColor=fills.get(ws[f"C{r}"].value, "FFFFFF"))
    sm = wb.create_sheet("Summary")
    for k in ("Present", "Review", "Absent", "NOT ENROLLED"):
        sm.append([k, sum(r[2] == k for r in rows)])
    sm.append(["Total", len(rows)]); sm.append(["Photos", len(paths)])
    sm.append(["match_thresh", args.match_thresh]); sm.append(["review_thresh", args.review_thresh])
    wb.save(os.path.join(out, "attendance.xlsx"))

    # ---------------------------------------------------------- validation report
    wb = Workbook(); ws = wb.active; ws.title = "Detections"
    write_sheet(ws, ["Class image", "Face #", "Detected face", "Matched enroll frame",
                     "Enroll Image1 (front)", "Pred Roll", "Pred Name", "Similarity",
                     "Margin top1-top2", "Decision", "Top1 roll (pre-assignment)",
                     "Face px", "Blur", "Det score", "Yaw ratio", "Feat norm"], face_rows,
                image_cols=("Detected face", "Matched enroll frame", "Enroll Image1 (front)"))
    rj = wb.create_sheet("Rejected (quality)")
    write_sheet(rj, ["Class image", "Face #", "Face", "Reasons", "Face px", "Blur", "Det score",
                     "Yaw ratio", "Feat norm"], rej_rows, image_cols=("Face",))
    wb.save(os.path.join(out, "validation_report.xlsx"))

    c = {k: sum(r[2] == k for r in rows) for k in fills}
    log.info("Present=%d Review=%d Absent=%d NotEnrolled=%d | faces kept=%d rejected=%d -> %s",
             c["Present"], c["Review"], c["Absent"], c["NOT ENROLLED"],
             len(face_rows), len(rej_rows), out)


if __name__ == "__main__":
    main()
